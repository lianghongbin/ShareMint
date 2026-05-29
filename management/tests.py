from decimal import Decimal
from io import BytesIO

from django.test import TestCase
from django.urls import reverse

from authentication.models import Role, User
from management.models import AuditLog, Investment, Profile
from management.services.member_import_export import (
    ADMIN_HEADERS,
    HEADMAN_HEADERS,
    build_admin_template_workbook,
    calc_gdt_quantity,
    export_admin_workbook,
    import_admin_workbook,
    import_headman_workbook,
)
from openpyxl import Workbook, load_workbook


class MemberImportExportTests(TestCase):
    def setUp(self):
        self.headman = User.objects.create_user(
            username='hm1',
            password='pass12345',
            role=Role.HEADMAN,
        )
        Profile.objects.create(user=self.headman, name='团长A', email='hm1@example.com')

    def test_calc_gdt_quantity_matches_frontend_formula(self):
        gdt = calc_gdt_quantity(Decimal('10000'), Decimal('5'), Decimal('1'))
        self.assertEqual(gdt, Decimal('9500'))

    def test_admin_template_has_expected_headers(self):
        workbook = load_workbook(BytesIO(build_admin_template_workbook()), read_only=True)
        headers = [cell.value for cell in next(workbook.active.iter_rows(min_row=1, max_row=1))]
        self.assertEqual(headers, ADMIN_HEADERS)
        workbook.close()

    def test_import_admin_creates_headman_only(self):
        wb = Workbook()
        ws = wb.active
        ws.append(ADMIN_HEADERS)
        ws.append(['new_headman', '新团长', '', 'hm-new@example.com', 'ShareMint123'])
        buffer = BytesIO()
        wb.save(buffer)

        result = import_admin_workbook(BytesIO(buffer.getvalue()))
        self.assertEqual(result.created_headmen, 1)
        self.assertEqual(result.created_members, 0)
        self.assertFalse(result.errors)
        self.assertTrue(User.objects.filter(username='new_headman', role=Role.HEADMAN).exists())

    def test_import_headman_creates_member_with_investment(self):
        wb = Workbook()
        ws = wb.active
        ws.append(HEADMAN_HEADERS)
        ws.append(['new_member', '新成员', '', '', '0x742d35Cc6634C0532925a3b844Bc9e7595f0bEb0', 'ShareMint123', 10000, 5, ''])
        buffer = BytesIO()
        wb.save(buffer)

        result = import_headman_workbook(self.headman, BytesIO(buffer.getvalue()))
        self.assertEqual(result.created_members, 1)
        self.assertFalse(result.errors)

        member = User.objects.get(username='new_member')
        self.assertEqual(member.referrer_id, self.headman.pk)
        investment = Investment.objects.get(user=member)
        self.assertEqual(investment.investment_amount, Decimal('10000'))
        self.assertEqual(investment.holding_quantity, Decimal('9500'))
        profile = Profile.objects.get(user=member)
        self.assertEqual(profile.bnb_wallet_address, '0x742d35cc6634c0532925a3b844bc9e7595f0beb0')

    def test_import_headman_rejects_invalid_bnb_wallet(self):
        wb = Workbook()
        ws = wb.active
        ws.append(HEADMAN_HEADERS)
        ws.append(['bad_wallet', '无效地址', '', '', 'not-a-wallet', 'ShareMint123', 1000, '0', ''])
        buffer = BytesIO()
        wb.save(buffer)

        result = import_headman_workbook(self.headman, BytesIO(buffer.getvalue()))
        self.assertEqual(result.created_members, 0)
        self.assertTrue(any('BNB' in err.message for err in result.errors))

    def test_import_headman_rejects_duplicate_username(self):
        User.objects.create_user(
            username='exists',
            password='pass12345',
            role=Role.MEMBER,
            referrer=self.headman,
        )
        wb = Workbook()
        ws = wb.active
        ws.append(HEADMAN_HEADERS)
        ws.append(['exists', '重复用户', '', '', '0x742d35Cc6634C0532925a3b844Bc9e7595f0bEb0', 'ShareMint123', '1000', '0', ''])
        buffer = BytesIO()
        wb.save(buffer)

        result = import_headman_workbook(self.headman, BytesIO(buffer.getvalue()))
        self.assertEqual(result.created_members, 0)
        self.assertTrue(any('已存在' in err.message for err in result.errors))

    def test_export_admin_includes_headmen_only(self):
        User.objects.create_user(
            username='mem1',
            password='pass12345',
            role=Role.MEMBER,
            referrer=self.headman,
        )

        workbook = load_workbook(BytesIO(export_admin_workbook()), read_only=True)
        rows = list(workbook.active.iter_rows(min_row=2, values_only=True))
        workbook.close()
        usernames = [row[0] for row in rows]
        self.assertIn('hm1', usernames)
        self.assertNotIn('mem1', usernames)


class AuditLogClearTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(
            username='admin_clear',
            password='pass12345',
            role=Role.ADMIN,
            is_staff=True,
            is_first_login=False,
        )
        Profile.objects.create(user=self.admin, name='Admin')
        AuditLog.objects.create(
            username='tester',
            role='admin',
            action='test:action',
            summary='测试操作',
            method='POST',
            path='/test/',
        )

    def test_admin_can_clear_audit_logs(self):
        self.client.force_login(self.admin)
        response = self.client.post(
            reverse('management:admin_settings_audit_clear'),
            {'confirmed': '1'},
        )
        self.assertRedirects(response, reverse('management:admin_settings_audit'))
        self.assertEqual(AuditLog.objects.count(), 1)
        log = AuditLog.objects.get()
        self.assertEqual(log.summary, '清空操作日志')

    def test_headman_cannot_clear_audit_logs(self):
        headman = User.objects.create_user(
            username='hm_clear',
            password='pass12345',
            role=Role.HEADMAN,
        )
        self.client.force_login(headman)
        response = self.client.post(
            reverse('management:admin_settings_audit_clear'),
            {'confirmed': '1'},
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(AuditLog.objects.count(), 1)
