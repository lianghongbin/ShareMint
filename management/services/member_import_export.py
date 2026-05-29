from __future__ import annotations

from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
from typing import BinaryIO

from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from authentication.email_change import validate_new_email
from authentication.models import Role, User
from authentication.users import is_username_taken
from commissions.models import SystemConfig
from management.models import Investment, Profile
from management.wallet import canonical_bnb_wallet_address, validate_bnb_wallet_address

DEFAULT_IMPORT_PASSWORD = 'ShareMint123'

ADMIN_HEADERS = [
    '用户名',
    '姓名',
    '手机号',
    '邮箱',
    '初始密码',
]

HEADMAN_HEADERS = [
    '用户名',
    '姓名',
    '手机号',
    '邮箱',
    'BNB钱包地址',
    '初始密码',
    '投资金额(USDT)',
    '手续费比例(%)',
    'GDT数量',
]

HEADER_FILL = PatternFill('solid', fgColor='1E293B')
HEADER_FONT = Font(bold=True, color='FFFFFF')


@dataclass
class ImportRowError:
    row_number: int
    message: str


@dataclass
class ImportResult:
    created_headmen: int = 0
    created_members: int = 0
    skipped: int = 0
    errors: list[ImportRowError] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.errors


def calc_gdt_quantity(
    amount: Decimal,
    fee_percent: Decimal,
    gdt_price: Decimal,
) -> Decimal | None:
    if amount <= 0 or gdt_price <= 0:
        return None
    if fee_percent < 0 or fee_percent >= 100:
        return None
    net = amount * (Decimal('1') - fee_percent / Decimal('100'))
    if net <= 0:
        return None
    return (net / gdt_price).quantize(Decimal('1'), rounding=ROUND_HALF_UP)


def estimate_fee_percent(
    amount: Decimal,
    holding: Decimal,
    gdt_price: Decimal,
) -> Decimal | None:
    if amount <= 0 or holding <= 0 or gdt_price <= 0:
        return None
    net = holding * gdt_price
    if net >= amount:
        return Decimal('0')
    fee = (Decimal('1') - net / amount) * Decimal('100')
    return fee.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)


def _parse_decimal(value, *, allow_none: bool = True) -> Decimal | None:
    text = str(value or '').strip().replace(',', '')
    if not text:
        return None if allow_none else Decimal('0')
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def _parse_row_values(row: tuple, headers: list[str]) -> dict[str, str]:
    data: dict[str, str] = {}
    for index, header in enumerate(headers):
        cell = row[index] if index < len(row) else None
        data[header] = str(cell).strip() if cell is not None else ''
    return data


def _is_empty_row(data: dict[str, str], required_keys: list[str]) -> bool:
    return not any(data.get(key) for key in required_keys)


def _style_header_row(sheet, headers: list[str]) -> None:
    sheet.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = sheet.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        sheet.column_dimensions[get_column_letter(col_idx)].width = 18


def _workbook_to_bytes(workbook: Workbook) -> bytes:
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_admin_template_workbook() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '导入模板'
    _style_header_row(sheet, ADMIN_HEADERS)
    sheet.append(['headman_demo', '示例团长', '13800000001', 'headman@example.com', DEFAULT_IMPORT_PASSWORD])
    sheet.append(['headman_demo2', '示例团长B', '13800000002', 'headman2@example.com', ''])
    notes = workbook.create_sheet('填写说明')
    notes.append(['说明'])
    notes.append(['1. 每行一名团长；团长不参与投资，仅负责发展会员。'])
    notes.append(['2. 会员请由对应团长在其账号下导入。'])
    notes.append(['3. 初始密码留空时默认使用 ShareMint123；导入后请提醒用户修改密码。'])
    notes.append(['4. 用户名不可重复；邮箱若已被其他账号绑定将无法导入。'])
    return _workbook_to_bytes(workbook)


def build_headman_template_workbook() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '导入模板'
    _style_header_row(sheet, HEADMAN_HEADERS)
    sheet.append(['member_demo1', '示例成员A', '13800000002', 'member1@example.com', '0x742d35Cc6634C0532925a3b844Bc9e7595f0bEb0', DEFAULT_IMPORT_PASSWORD, '10000', '5', ''])
    sheet.append(['member_demo2', '示例成员B', '13800000003', 'member2@example.com', '0x742d35Cc6634C0532925a3b844Bc9e7595f0bEb1', DEFAULT_IMPORT_PASSWORD, '5000', '0', ''])
    notes = workbook.create_sheet('填写说明')
    notes.append(['说明'])
    notes.append(['1. 每行一名成员；投资金额与手续费比例用于计算 GDT 持有数量。'])
    notes.append(['2. BNB钱包地址必填，须为 BNB Smart Chain（BEP20）链地址（0x 开头，42 位）；填错链或填错地址将导致转入资金无法找回。'])
    notes.append(['3. GDT 数量可留空，系统按当前 GDT 价格自动计算。'])
    notes.append(['4. 初始密码留空时默认使用 ShareMint123。'])
    return _workbook_to_bytes(workbook)


def export_admin_workbook() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '团长列表'
    _style_header_row(sheet, ADMIN_HEADERS)

    headmen = (
        User.objects.filter(role=Role.HEADMAN)
        .select_related('profile')
        .order_by('username')
    )
    for headman in headmen:
        profile = getattr(headman, 'profile', None)
        sheet.append([
            headman.username,
            profile.name if profile else headman.username,
            profile.phone if profile else '',
            profile.email if profile else '',
            '',
        ])
    return _workbook_to_bytes(workbook)


def export_headman_workbook(headman: User) -> bytes:
    gdt_price = SystemConfig.get_gdt_price()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '团队成员'
    _style_header_row(sheet, HEADMAN_HEADERS)

    members = (
        headman.referred_members.filter(role=Role.MEMBER)
        .select_related('profile')
        .prefetch_related('investments')
        .order_by('username')
    )
    for member in members:
        member_profile = getattr(member, 'profile', None)
        investments = member.investments.all()
        total_amount = sum((inv.investment_amount for inv in investments), Decimal('0'))
        total_gdt = sum((inv.holding_quantity for inv in investments), Decimal('0'))
        fee_percent = estimate_fee_percent(total_amount, total_gdt, gdt_price)
        sheet.append([
            member.username,
            member_profile.name if member_profile else member.username,
            member_profile.phone if member_profile else '',
            member_profile.email if member_profile else '',
            member_profile.bnb_wallet_address if member_profile else '',
            '',
            float(total_amount) if total_amount else '',
            float(fee_percent) if fee_percent is not None else '',
            float(total_gdt) if total_gdt else '',
        ])
    return _workbook_to_bytes(workbook)


def _resolve_gdt_quantity(
    amount: Decimal | None,
    fee_percent: Decimal,
    explicit_gdt: Decimal | None,
    gdt_price: Decimal,
    row_number: int,
    errors: list[ImportRowError],
) -> Decimal | None:
    if explicit_gdt is not None:
        if explicit_gdt < 0:
            errors.append(ImportRowError(row_number, 'GDT 数量不能为负数。'))
            return None
        return explicit_gdt
    if amount is None or amount <= 0:
        return None
    gdt = calc_gdt_quantity(amount, fee_percent, gdt_price)
    if gdt is None:
        errors.append(ImportRowError(row_number, '无法根据投资金额与手续费比例计算 GDT 数量。'))
    return gdt


def _create_headman_user(
    *,
    username: str,
    password: str,
    name: str,
    phone: str,
    email: str,
) -> User:
    user = User.objects.create_user(
        username=username,
        password=password,
        role=Role.HEADMAN,
        is_first_login=True,
    )
    if email:
        user.email = email
        user.save(update_fields=['email'])
    Profile.objects.create(
        user=user,
        name=name,
        phone=phone,
        email=email,
    )
    return user


def _create_member_user(
    *,
    headman: User,
    username: str,
    password: str,
    name: str,
    phone: str,
    email: str,
    bnb_wallet_address: str,
    investment_amount: Decimal | None,
    holding_quantity: Decimal | None,
) -> User:
    member = User.objects.create_user(
        username=username,
        password=password,
        role=Role.MEMBER,
        referrer=headman,
        is_first_login=True,
    )
    if email:
        member.email = email
        member.save(update_fields=['email'])
    Profile.objects.create(
        user=member,
        name=name,
        phone=phone,
        email=email,
        bnb_wallet_address=bnb_wallet_address,
    )
    if investment_amount and investment_amount > 0 and holding_quantity and holding_quantity > 0:
        Investment.objects.create(
            user=member,
            investment_date=timezone.now(),
            investment_amount=investment_amount,
            holding_quantity=holding_quantity,
        )
    return member


def import_admin_workbook(file: BinaryIO) -> ImportResult:
    workbook = load_workbook(file, read_only=True, data_only=True)
    sheet = workbook.active
    result = ImportResult()

    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        data = _parse_row_values(row, ADMIN_HEADERS)
        if _is_empty_row(data, ['用户名', '姓名']):
            continue

        username = data['用户名']
        name = data['姓名']
        phone = data['手机号']
        email = data['邮箱']
        password = data['初始密码'] or DEFAULT_IMPORT_PASSWORD

        if not username:
            result.errors.append(ImportRowError(row_number, '用户名不能为空。'))
            continue
        if not name:
            result.errors.append(ImportRowError(row_number, '姓名不能为空。'))
            continue
        if len(password) < 8:
            result.errors.append(ImportRowError(row_number, '初始密码至少 8 位。'))
            continue
        if is_username_taken(username):
            result.errors.append(ImportRowError(row_number, f'用户名 {username} 已存在。'))
            continue
        if email:
            email_error = validate_new_email(User(), email)
            if email_error:
                result.errors.append(ImportRowError(row_number, email_error))
                continue
        try:
            with transaction.atomic():
                _create_headman_user(
                    username=username,
                    password=password,
                    name=name,
                    phone=phone,
                    email=email,
                )
            result.created_headmen += 1
        except Exception as exc:
            result.errors.append(ImportRowError(row_number, f'创建团长失败：{exc}'))

    workbook.close()
    return result


def import_headman_workbook(headman: User, file: BinaryIO) -> ImportResult:
    workbook = load_workbook(file, read_only=True, data_only=True)
    sheet = workbook.active
    gdt_price = SystemConfig.get_gdt_price()
    result = ImportResult()

    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        data = _parse_row_values(row, HEADMAN_HEADERS)
        if _is_empty_row(data, ['用户名', '姓名']):
            continue

        username = data['用户名']
        name = data['姓名']
        phone = data['手机号']
        email = data['邮箱']
        bnb_wallet_raw = data['BNB钱包地址']
        password = data['初始密码'] or DEFAULT_IMPORT_PASSWORD

        if not username:
            result.errors.append(ImportRowError(row_number, '用户名不能为空。'))
            continue
        if not name:
            result.errors.append(ImportRowError(row_number, '姓名不能为空。'))
            continue
        if len(password) < 8:
            result.errors.append(ImportRowError(row_number, '初始密码至少 8 位。'))
            continue
        if is_username_taken(username):
            result.errors.append(ImportRowError(row_number, f'用户名 {username} 已存在。'))
            continue
        if email:
            email_error = validate_new_email(User(), email)
            if email_error:
                result.errors.append(ImportRowError(row_number, email_error))
                continue
        wallet_error = validate_bnb_wallet_address(bnb_wallet_raw)
        if wallet_error:
            result.errors.append(ImportRowError(row_number, wallet_error))
            continue
        bnb_wallet_address = canonical_bnb_wallet_address(bnb_wallet_raw)

        amount = _parse_decimal(data['投资金额(USDT)'])
        fee_percent = _parse_decimal(data['手续费比例(%)'], allow_none=True) or Decimal('0')
        explicit_gdt = _parse_decimal(data['GDT数量'])
        if fee_percent < 0 or fee_percent >= 100:
            result.errors.append(ImportRowError(row_number, '手续费比例须为 0 到 100 之间的数字。'))
            continue
        holding = _resolve_gdt_quantity(amount, fee_percent, explicit_gdt, gdt_price, row_number, result.errors)
        if holding is None and amount and amount > 0 and explicit_gdt is None:
            continue

        try:
            with transaction.atomic():
                _create_member_user(
                    headman=headman,
                    username=username,
                    password=password,
                    name=name,
                    phone=phone,
                    email=email,
                    bnb_wallet_address=bnb_wallet_address,
                    investment_amount=amount,
                    holding_quantity=holding,
                )
            result.created_members += 1
        except Exception as exc:
            result.errors.append(ImportRowError(row_number, f'创建成员失败：{exc}'))

    workbook.close()
    return result
